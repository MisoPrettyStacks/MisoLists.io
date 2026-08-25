#!/usr/bin/env python3
"""
Strategic Investment Pipeline — Autonomous Ingestion Engine (GitHub Pages edition)
===============================================================================
Runs locally and in GitHub Actions. Reads every file in ./reports, SHA-256
dedupes, parses xlsx + pdf, merges the two tracker views by company, rebuilds
the SQLite database (data.db) with full source lineage, and regenerates the
dashboard data bundle (assets/data.js). GitHub Pages auto-rebuilds from the
committed assets/data.js.

Run:
  python3 ingest.py            # default: ./reports  (repo-relative)
  python3 ingest.py <folder>   # any folder of xlsx/pdf reports

Triggered automatically by .github/workflows/ingest.yml on every push to
reports/** and on a daily schedule — fully free on GitHub Actions.
"""
import openpyxl, json, os, sqlite3, hashlib, sys, glob, re
from datetime import datetime, timezone

BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(BASE, "data")
SITE_DATA = os.path.join(BASE, "assets", "data.js")
DB_PATH = os.path.join(BASE, "data.db")
DEFAULT_FOLDER = os.path.join(BASE, "reports")

def now(): return datetime.now(timezone.utc).isoformat()

def sha(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()

def sheet_rows(ws):
    return [[c for c in r] for r in ws.iter_rows(values_only=True)]

def load_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    return {sn: sheet_rows(wb[sn]) for sn in wb.sheetnames}

def load_pdf(path):
    try:
        import pdfplumber
        pages = []
        with pdfplumber.open(path) as pdf:
            for pg in pdf.pages:
                pages.append(pg.extract_text() or "")
        return {"pages": len(pages), "text": "\n\n".join(pages)}
    except Exception as e:
        return {"pages": 0, "text": "", "error": str(e)}

def num(v):
    if v is None or v == "": return None
    try: return float(v)
    except: return None

def prob_num(p):
    if not p: return None
    m = re.match(r"(\d+(?:\.\d+)?)", str(p))
    return float(m.group(1)) if m else None

def merge_trackers(colored_rowsets, graded_rowsets):
    """colored_rowsets / graded_rowsets are LISTS of row-tables — one per
    report file that contains a matching sheet — so that dropping in
    additional "Master Pipeline Tracker" / "Graded Pipeline" report files
    unions them together instead of the ingester picking only one file
    and silently ignoring the rest."""
    gl = {}
    for graded_rows in graded_rowsets:
        for r in (graded_rows or [])[1:]:
            if r and r[0]: gl[r[0].strip().lower()] = r

    merged, seen = [], set()
    for colored_rows in colored_rowsets:
        for r in (colored_rows or [])[1:]:
            if not r or not r[2]: continue
            company = r[2]
            key = company.strip().lower()
            if key in seen: continue  # first occurrence wins (report order)
            seen.add(key)
            g = gl.get(key)
            def gs(i): return g[i] if g and i < len(g) else None
            merged.append({
                "company": company, "primary_event": r[0], "thematic_probability": r[1],
                "category": r[3], "funding_stage": r[4], "ipo_date": r[5],
                "sector": r[6] or gs(4), "source": r[7] or gs(5), "notes": r[8], "historical_flag": r[9],
                "stage_score": num(gs(6)), "financial_score": num(gs(7)), "capital_score": num(gs(8)),
                "market_score": num(gs(9)), "risk_adj": num(gs(10)), "total_score": num(gs(11)), "grade": gs(12),
                "stage_rationale": gs(13), "financial_rationale": gs(14), "capital_rationale": gs(15),
                "market_rationale": gs(16), "risk_rationale": gs(17),
            })
    merged.sort(key=lambda x: (x.get("total_score") or 0), reverse=True)
    return merged

def main(folder=DEFAULT_FOLDER):
    os.makedirs(DATA, exist_ok=True)
    ts = now()
    files = sorted(glob.glob(os.path.join(folder, "*")))
    if not files:
        print(f"[ingest] No files found in {folder}"); return

    db = sqlite3.connect(DB_PATH)
    c = db.cursor()
    c.executescript("""
      CREATE TABLE IF NOT EXISTS files (id INTEGER PRIMARY KEY, name TEXT, sha256 TEXT, size INT, ingested_at TEXT, sheets TEXT, row_count INT);
      CREATE TABLE IF NOT EXISTS pipeline (company TEXT, primary_event TEXT, thematic_probability TEXT, category TEXT, funding_stage TEXT,
        ipo_date TEXT, sector TEXT, source TEXT, notes TEXT, historical_flag TEXT, stage_score REAL, financial_score REAL, capital_score REAL,
        market_score REAL, risk_adj REAL, total_score REAL, grade TEXT, stage_rationale TEXT, financial_rationale TEXT, capital_rationale TEXT,
        market_rationale TEXT, risk_rationale TEXT, ingested_at TEXT);
      CREATE TABLE IF NOT EXISTS methodology (file TEXT, line TEXT);
    """)
    existing = {r[0] for r in c.execute("SELECT sha256 FROM files")}
    lineage, payloads, new_count = [], {}, 0

    for f in files:
        if os.path.isdir(f): continue
        name = os.path.basename(f)
        h = sha(f)
        is_new = h not in existing
        if is_new:
            new_count += 1
            print(f"[ingest] INGEST (new): {name}  (sha {h[:16]})")
            c.execute("INSERT INTO files(name,sha256,size,ingested_at) VALUES(?,?,?,?)", (name, h, os.path.getsize(f), ts))
            existing.add(h); ing_ts = ts
        else:
            print(f"[ingest] already ingested (re-parsing for rebuild): {name}")
            ing_ts = (c.execute("SELECT ingested_at FROM files WHERE sha256=?", (h,)).fetchone() or [ts])[0]
        lineage.append({"name": name, "sha256": h[:16], "size": os.path.getsize(f), "ingested_at": ing_ts})
        if name.lower().endswith(".xlsx"): payloads[name] = load_xlsx(f)
        elif name.lower().endswith(".pdf"): payloads[name] = load_pdf(f)

    def find_xlsx(pred):
        for n, d in payloads.items():
            if n.lower().endswith(".xlsx") and pred(d): return d
        return None

    def find_all_xlsx(pred):
        return [d for n, d in payloads.items() if n.lower().endswith(".xlsx") and pred(d)]

    col_all = find_all_xlsx(lambda d: any("Master Pipeline Tracker" in s for s in d))
    grd_all = find_all_xlsx(lambda d: any("Graded Pipeline" in s for s in d))
    orgs = find_xlsx(lambda d: any("Master Overview" in s for s in d))
    icf = find_xlsx(lambda d: any("Comprehensive Report" in s for s in d))
    pdf = next((d for n, d in payloads.items() if n.lower().endswith(".pdf")), None)

    c.execute("DELETE FROM pipeline"); c.execute("DELETE FROM methodology")
    colored_rowsets = [d.get("Master Pipeline Tracker") for d in col_all]
    colored_sum = next((d.get("Summary & Methodology") for d in col_all if d.get("Summary & Methodology")), None)
    graded_rowsets = [d.get("Graded Pipeline") for d in grd_all]
    graded_meth = next((d.get("Methodology") for d in grd_all if d.get("Methodology")), None)
    orgs_master = orgs.get("Master Overview") if orgs else None
    orgs_readme = orgs.get("README") if orgs else None

    merged = merge_trackers(colored_rowsets, graded_rowsets) if colored_rowsets else []
    for m in merged:
        c.execute("INSERT INTO pipeline (" + ",".join(m.keys()) + ",ingested_at) VALUES (" + ",".join(["?"] * len(m)) + ",?)",
                  list(m.values()) + [ts])
    if graded_meth:
        for ln in graded_meth:
            if ln and ln[0]: c.execute("INSERT INTO methodology(file,line) VALUES(?,?)", ("GRADED", ln[0]))
    db.commit()

    meth_lines = [m[0] for m in graded_meth if m and m[0]] if graded_meth else []
    first_colored = next((r for r in colored_rowsets if r), [[]])
    first_graded = next((r for r in graded_rowsets if r), [[]])
    headers = {"colored": first_colored[0] if first_colored else [], "graded": first_graded[0] if first_graded else []}
    payload = {
        "generated_at": ts, "headers": headers, "companies": merged, "count": len(merged),
        "methodology": meth_lines, "colored_summary": colored_sum or [],
        "orgs_overview": orgs_master or [], "orgs_readme": [r[0] for r in orgs_readme if r and r[0]] if orgs_readme else [],
        "ic_sheets": list(icf.keys()) if icf else [], "ic_comprehensive": icf.get("Comprehensive Report", []) if icf else [],
        "ic_analysis": icf.get("Analysis (Cross-Tally)", []) if icf else [],
        "pdf_pages": pdf["pages"] if pdf else 0, "pdf_text": pdf["text"] if pdf else "",
        "lineage": lineage,
    }
    os.makedirs(os.path.dirname(SITE_DATA), exist_ok=True)
    with open(SITE_DATA, "w") as fh:
        fh.write("window.PIPELINE_DATA = " + json.dumps(payload, default=str) + ";")

    print(f"\n[ingest] DONE. New files: {new_count} | Entities: {len(merged)} | Sources: {len(lineage)}")
    print(f"[ingest] Database: {DB_PATH}")
    print(f"[ingest] Dashboard data: {SITE_DATA}")

if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_FOLDER)
