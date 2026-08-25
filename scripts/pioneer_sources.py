#!/usr/bin/env python3
"""
Pioneer & Innovator Source Scraper — Autonomous Discovery Engine
==================================================================
Agentically fetches a registry of international tech-pioneer / innovator
recognition programs, extracts named entities where the source page's
structure allows reliable extraction, and writes a report in the exact
Master Pipeline Tracker / Graded Pipeline schema that ingest.py already
consumes. Drop the output into reports/ (this script does that itself)
and the existing ingest.yml workflow (push to reports/**, or its daily
cron) will merge it into the live dashboard automatically. No manual
spreadsheet work required after the first run.

Design principle: this script never invents financial/funding facts.
Every entity it emits is scored using ONLY what the methodology in
Master_Investment_Pipeline_Tracker_GRADED already defines as knowable
from a bare cohort/award selection — Stage=5 (private/unknown, award-
selection only), Financial=0, Capital=0, Market=0, Risk=-5 (no
disclosed liquidity timeline) — the same convention already used in
this repo for NATO DIANA / G20 TechSprint cohort-only rows. If a
source can't be reliably parsed (JS-rendered SPA, blocked, empty),
it is logged as NEEDS_REVIEW and NOT silently skipped or faked.

Run:
  python3 scripts/pioneer_sources.py            # scrape all, write report
  python3 scripts/pioneer_sources.py --dry-run  # scrape + print, no file
"""
import argparse
import json
import os
import re
import sys
from datetime import datetime, timezone

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font

try:
    from playwright.sync_api import sync_playwright
    HAVE_PLAYWRIGHT = True
except ImportError:
    HAVE_PLAYWRIGHT = False

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPORTS_DIR = os.path.join(BASE, "reports")
SEED_DIR = os.path.join(BASE, "scripts", "seed_data")
UA = "Mozilla/5.0 (compatible; MisoListsBot/1.0; +https://github.com/MisoPrettyStacks/MisoLists.io)"
TIMEOUT = 20

# ---------------------------------------------------------------------------
# Source registry (post Global-Media-Rankings-section removal per request)
# ---------------------------------------------------------------------------
SOURCES = [
    # -- World Economic Forum (4 lists, one pipeline) --
    {"id": "wef_tech_pioneers", "name": "WEF Technology Pioneers",
     "url": "https://initiatives.weforum.org/technology-pioneers/home",
     "category": "WEF", "js_rendered": True},
    {"id": "wef_global_innovators", "name": "WEF Global Innovators",
     "url": "https://initiatives.weforum.org/innovator-communities/globalinnovators",
     "category": "WEF", "js_rendered": True},
    {"id": "wef_unicorns", "name": "WEF Unicorns Community",
     "url": "https://initiatives.weforum.org/innovator-communities/organizations",
     "category": "WEF", "js_rendered": True},
    {"id": "wef_uplink", "name": "WEF UpLink Innovation Challenges",
     "url": "https://uplink.weforum.org/uplink-innovation-challenge-series-spring-2026",
     "category": "WEF", "js_rendered": True},
    # -- UN / Intergovernmental --
    {"id": "un_itu_wsis", "name": "UN ITU WSIS Prizes",
     "url": "https://www.itu.int/en/mediacentre/Pages/PR-2026-07-09-WSIS-Prizes.aspx",
     "category": "UN/IGO", "js_rendered": False},
    {"id": "worldbank_govtech", "name": "World Bank GovTech Innovation Challenge Awards",
     "url": "https://www.worldbank.org/en/events/2025/11/26/govtech-innovation-challenge-award",
     "category": "UN/IGO", "js_rendered": False},
    {"id": "wipo_gii", "name": "WIPO Global Innovation Index",
     "url": "https://www.wipo.int/en/web/global-innovation-index",
     "category": "UN/IGO", "js_rendered": False},
    {"id": "bis_innovation_hub", "name": "BIS Innovation Hub Project Portfolio",
     "url": "https://www.bis.org/about/bisih/about.htm",
     "category": "UN/IGO", "js_rendered": False},
    # -- Global Ecosystem / Ranking Bodies --
    {"id": "startupblink", "name": "StartupBlink Global Startup Ecosystem Index",
     "url": "https://lp.startupblink.com/report/",
     "category": "Ecosystem Index", "js_rendered": True},
    {"id": "startup_genome", "name": "Startup Genome Global Startup Ecosystem Report",
     "url": "https://startupgenome.com/report/the-global-startup-ecosystem-report-2026/",
     "category": "Ecosystem Index", "js_rendered": True},
    {"id": "dealroom_index", "name": "Dealroom Global Tech Ecosystem Index",
     "url": "https://dealroom.co/tech-ecosystem-index-2026/",
     "category": "Ecosystem Index", "js_rendered": True},
    # -- Global Non-Profit / Network Organizations --
    {"id": "endeavor_outliers", "name": "Endeavor Global Entrepreneurs / Outliers",
     "url": "https://endeavor.org/2026-endeavor-outliers/",
     "category": "Endeavor", "js_rendered": False},
    {"id": "endeavor_portfolio", "name": "Endeavor Entrepreneur Companies Portfolio",
     "url": "https://endeavor.org/entrepreneur-companies/",
     "category": "Endeavor", "js_rendered": True},
    # -- Global Pitch Competitions / Convening Platforms --
    {"id": "slush_100", "name": "Slush 100",
     "url": "https://www.slush.org",
     "category": "Pitch Competition", "js_rendered": True},
    {"id": "startup_world_cup", "name": "Startup World Cup",
     "url": "https://www.startupworldcup.io",
     "category": "Pitch Competition", "js_rendered": True},
    {"id": "gitex_supernova", "name": "GITEX / Expand North Star Supernova Challenge",
     "url": "https://www.expandnorthstar.com",
     "category": "Pitch Competition", "js_rendered": True},
]

# Words that show up in nav/footer/legal boilerplate — filtered out of
# candidate entity names so we don't pollute the pipeline with junk like
# "Privacy Policy" or "Sign Up".
STOPWORDS = {
    "home", "about", "contact", "privacy", "policy", "terms", "sign in",
    "sign up", "log in", "cookie", "cookies", "menu", "search", "subscribe",
    "newsletter", "read more", "learn more", "explore", "download", "share",
    "follow us", "all rights reserved", "skip to content", "back to top",
}


def now():
    return datetime.now(timezone.utc).isoformat()


BROWSER_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36")


def fetch_rendered(url):
    """Fetch a page with a real headless browser: handles client-side
    rendered SPAs and most basic bot-detection walls that a plain
    requests.get cannot get past. Requires `playwright install chromium`
    to have been run (see .github/workflows/pioneer-sources.yml)."""
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            page = browser.new_page(user_agent=BROWSER_UA, viewport={"width": 1366, "height": 900})
            page.goto(url, timeout=30000, wait_until="networkidle")
            page.wait_for_timeout(1500)  # let lazy-loaded lists settle
            html = page.content()
            browser.close()
            return html
    except Exception as e:
        return None, f"playwright error: {e}"


def fetch_static(url):
    try:
        r = requests.get(url, headers={
            "User-Agent": BROWSER_UA,
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "en-US,en;q=0.9",
        }, timeout=TIMEOUT)
        r.raise_for_status()
        return r.text
    except Exception as e:
        return None, str(e)





def looks_like_entity_name(text):
    t = text.strip()
    if not (2 <= len(t) <= 60):
        return False
    if t.lower() in STOPWORDS:
        return False
    if any(sw in t.lower() for sw in STOPWORDS):
        return False
    if t.isupper() and len(t) > 20:  # shouty boilerplate banners
        return False
    if not re.search(r"[A-Za-z]", t):
        return False
    if re.match(r"^\d+$", t):
        return False
    # crude "looks like a proper noun / brand" heuristic: starts with a
    # capital letter, isn't a full sentence (no terminal punctuation)
    if not t[0].isupper():
        return False
    if t.endswith((".", "!", "?")) and len(t.split()) > 6:
        return False
    return True


def extract_candidates(html):
    """Best-effort entity-name extraction from static HTML. Returns a
    deduped list. Deliberately conservative — false negatives (missing
    a real name) are fine, false positives (nav junk scored as a company)
    are not."""
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()

    candidates = []
    # Headings and list items are the highest-signal locations for
    # "winner"/"cohort" style pages.
    for tag in soup.find_all(["h2", "h3", "h4", "li", "strong", "b"]):
        text = tag.get_text(" ", strip=True)
        if looks_like_entity_name(text):
            candidates.append(text)

    seen, out = set(), []
    for c in candidates:
        key = c.lower()
        if key not in seen:
            seen.add(key)
            out.append(c)
    return out


def scrape_source(src):
    result = {
        "id": src["id"], "name": src["name"], "url": src["url"],
        "category": src["category"], "status": None, "entities": [],
        "fetched_at": now(),
    }
    method_notes = []
    candidates = []

    # Attempt 1: cheap static fetch, unless we already know it's a pure SPA.
    if not src["js_rendered"]:
        fetched = fetch_static(src["url"])
        if isinstance(fetched, tuple):
            method_notes.append(f"static fetch failed ({fetched[1]})")
        else:
            candidates = extract_candidates(fetched)
            method_notes.append(f"static fetch: {len(candidates)} candidates")

    # Attempt 2: headless browser, if attempt 1 didn't yield enough (or we
    # already knew this was a JS-rendered page) and Playwright is available.
    if len(candidates) < 3 and HAVE_PLAYWRIGHT:
        rendered = fetch_rendered(src["url"])
        if isinstance(rendered, str):
            browser_candidates = extract_candidates(rendered)
            method_notes.append(f"browser fetch: {len(browser_candidates)} candidates")
            if len(browser_candidates) > len(candidates):
                candidates = browser_candidates
        else:
            method_notes.append(f"browser fetch failed ({rendered[1] if isinstance(rendered, tuple) else 'unknown error'})")
    elif len(candidates) < 3 and not HAVE_PLAYWRIGHT:
        method_notes.append("Playwright not installed — cannot attempt browser render")

    if len(candidates) < 3:
        result["status"] = f"NEEDS_REVIEW: {'; '.join(method_notes)} — page likely doesn't list entities directly, or is blocked."
        result["entities"] = candidates
        return result

    result["status"] = f"OK: {len(candidates)} candidate entities extracted ({'; '.join(method_notes)}) — unverified, review before treating as ground truth"
    result["entities"] = candidates[:40]  # cap to avoid runaway noise
    return result


def load_seed(source_id):
    """Verified, manually-researched entities for sources where live
    scraping can't reach the actual roster (JS SPA, PDF-only release,
    etc). Stored as small JSON files in scripts/seed_data/ so they can
    be refreshed by hand or by a future research pass without touching
    this script."""
    path = os.path.join(SEED_DIR, f"{source_id}.json")
    if os.path.exists(path):
        with open(path) as f:
            return json.load(f)
    return []


def build_workbook(all_results, out_path):
    wb = openpyxl.Workbook()

    # -- Sheet 1: Master Pipeline Tracker (colored-workbook schema) --
    ws1 = wb.active
    ws1.title = "Master Pipeline Tracker"
    ws1.append(["Primary Event", "Thematic Probability", "Company or Entity Name",
                "Category", "Funding Stage", "IPO Date", "Sector or Theme",
                "Source Report or Organization", "Notes on Invention, Innovation, "
                "Product, or Strategic Relevance", "Historical Flag"])
    for c in ws1[1]:
        c.font = Font(bold=True)

    # -- Sheet 2: Graded Pipeline (graded-workbook schema) --
    wb.create_sheet("Graded Pipeline")
    ws2 = wb["Graded Pipeline"]
    ws2.append(["Company / Entity", "Category", "Funding Stage", "IPO Date", "Sector",
                "Source", "Stage Score (0-25)", "Financial Traction (0-30)",
                "Capital Backing (0-20)", "Market Position (0-15)",
                "Risk Adj. (-15 to 0)", "TOTAL SCORE (0-100)", "GRADE",
                "Stage Rationale", "Financial Rationale", "Capital Rationale",
                "Market Rationale", "Risk Rationale",
                "Historical Flag (WEF/pre-2020 alumni)"])
    for c in ws2[1]:
        c.font = Font(bold=True)

    # -- Sheet 3: Scrape Log (extra sheet, ignored by ingest.py, human-readable
    #    audit trail of what was auto-discovered vs. what needs a human) --
    wb.create_sheet("Scrape Log")
    ws3 = wb["Scrape Log"]
    ws3.append(["Source", "Category", "URL", "Status", "Fetched At (UTC)"])
    for c in ws3[1]:
        c.font = Font(bold=True)

    row_count = 0
    for res in all_results:
        ws3.append([res["name"], res["category"], res["url"], res["status"], res["fetched_at"]])
        for entity in res["entities"]:
            stage, fin, cap, mkt, risk = 5, 0, 0, 0, -5
            total = stage + fin + cap + mkt + risk
            grade = "A" if total >= 35 else "B" if total >= 25 else "C" if total >= 15 else "D" if total >= 8 else "F"
            source_label = f"{res['name']} (auto-discovered {res['fetched_at'][:10]})"
            notes = (f"Auto-discovered via agentic source scan of {res['name']}. "
                     f"UNVERIFIED — confirm entity identity and current status before acting on this row.")
            ws1.append(["", "N/A", entity, "Unknown", "Unknown", "", res["category"],
                        source_label, notes, "No"])
            ws2.append([entity, "Unknown", "Unknown", "", res["category"], source_label,
                        stage, fin, cap, mkt, risk, total, grade,
                        "Private/Unknown, award-cohort selection only, no disclosed funding stage",
                        "0 no financial performance data disclosed",
                        "0 no capital-backing detail disclosed",
                        "0 no market-position evidence disclosed (selection alone is not market-leadership evidence)",
                        "-5 no disclosed liquidity timeline (data opacity)",
                        "No"])
            row_count += 1

    wb.save(out_path)
    return row_count


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    os.makedirs(REPORTS_DIR, exist_ok=True)
    all_results = []
    for src in SOURCES:
        print(f"[pioneer_sources] scanning {src['name']} ...")
        res = scrape_source(src)
        if not res["entities"]:
            seeded = load_seed(src["id"])
            if seeded:
                res["entities"] = seeded
                res["status"] += f" | seeded fallback used ({len(seeded)} verified entities from scripts/seed_data/{src['id']}.json)"
        print(f"    -> {res['status']}")
        all_results.append(res)

    if args.dry_run:
        print(json.dumps(all_results, indent=2))
        return

    out_path = os.path.join(REPORTS_DIR, "Pioneer_Innovator_Sources_AUTO.xlsx")
    n = build_workbook(all_results, out_path)
    print(f"\n[pioneer_sources] wrote {out_path} — {n} entity rows across {len(SOURCES)} sources")
    print("[pioneer_sources] run `python3 ingest.py` (or push this file) to merge into the live dashboard")


if __name__ == "__main__":
    main()
